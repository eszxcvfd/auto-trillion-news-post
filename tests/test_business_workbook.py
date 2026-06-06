import os
import shutil
import tempfile
import unittest
import openpyxl

from src.models import NewsItem, BusinessWorkbookRow
from src.business_workbook import (
    normalize_header,
    normalize_cell_content,
    parse_row_id,
    detect_workbook_type,
    ingest_business_workbook,
    map_news_item_to_business_row,
    get_sheet_name_for_keyword,
    save_news_to_business_excel,
    generate_drafts_for_business_excel
)

class TestBusinessWorkbook(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.legacy_file = os.path.join(self.test_dir, "legacy_workbook.xlsx")
        self.business_file = os.path.join(self.test_dir, "business_workbook.xlsx")
        self.invalid_file = os.path.join(self.test_dir, "invalid_workbook.xlsx")
        
    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def test_normalize_header(self):
        self.assertEqual(normalize_header("  TikTok  "), "tiktok")
        self.assertEqual(normalize_header(" X (Twitter)\n"), "x (twitter)")
        self.assertEqual(normalize_header(None), None)
        self.assertEqual(normalize_header(123), "123")

    def test_normalize_cell_content(self):
        self.assertEqual(normalize_cell_content("   "), None)
        self.assertEqual(normalize_cell_content("."), None)
        self.assertEqual(normalize_cell_content("  actual text  "), "actual text")
        self.assertEqual(normalize_cell_content(None), None)

    def test_parse_row_id(self):
        self.assertEqual(parse_row_id(1.0), 1)
        self.assertEqual(parse_row_id(2), 2)
        self.assertEqual(parse_row_id(" 3 "), 3)
        self.assertEqual(parse_row_id("abc"), "abc")
        self.assertEqual(parse_row_id(None), None)

    def test_detect_workbook_type(self):
        # 1. Create a legacy workbook (Contract A)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Trillion News"
        legacy_headers = [
            "ID", "Found Date", "Keyword", "Title", "Source", "URL", 
            "Snippet", "Published Text", "Image File", "Platform", 
            "Top Hashtags", "Generated Post File", "Status", "Notes"
        ]
        ws.append(legacy_headers)
        wb.save(self.legacy_file)
        
        # 2. Create a business workbook (Contract B)
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = "Payment"
        business_headers = [
            "#", "Trillion $ news Title", "Image link", "Linkedin", 
            "Facebook", "X (Twitter)", "Instagram", "Pinterest", 
            "Threads", "TikTok", "YouTube", "Link Post"
        ]
        ws2.append(business_headers)
        wb2.save(self.business_file)
        
        # 3. Create an invalid/unrecognized workbook
        wb3 = openpyxl.Workbook()
        ws3 = wb3.active
        ws3.title = "Sheet1"
        ws3.append(["Random Column", "Another Column"])
        wb3.save(self.invalid_file)
        
        # Verify detection
        self.assertEqual(detect_workbook_type(self.legacy_file), "legacy")
        self.assertEqual(detect_workbook_type(self.business_file), "business")
        
        with self.assertRaises(ValueError):
            detect_workbook_type(self.invalid_file)
            
        with self.assertRaises(FileNotFoundError):
            detect_workbook_type("nonexistent_file.xlsx")

    def test_ingest_business_workbook_success(self):
        # Create a valid Business Workbook with two sheets
        wb = openpyxl.Workbook()
        
        # Sheet 1: Payment (with whitespace in headers, empty draft cells, and '.')
        ws1 = wb.active
        ws1.title = "Payment"
        ws1.append([
            " # ", "Trillion $ news Title ", "Image link", "Linkedin", 
            "Facebook", " X (Twitter) ", "Instagram", "Pinterest", 
            "Threads", "TikTok ", "YouTube", "Link Post"
        ])
        # Add some rows
        # Row 2: fully valid
        ws1.append([
            1.0, 
            "Merchant Payments trillion $ Opportunity", 
            "https://drive.google.com/file/123", 
            "LinkedIn draft content", 
            "Facebook draft content", 
            " X draft content ", 
            ".", # Instagram (should be normalized to None/skipped)
            " ", # Pinterest (should be normalized to None/skipped)
            "Threads draft content",
            ".", # TikTok (should be None)
            "YouTube draft content",
            "LinkedIn: https://linkedin.com/post1\nFacebook: https://facebook.com/post1"
        ])
        # Row 3: missing title (should be skipped)
        ws1.append([
            2.0, 
            " ", # Empty title
            "https://drive.google.com/file/456", 
            "LinkedIn draft 2", "Facebook draft 2", "X draft 2", "Instagram draft 2",
            "Pinterest draft 2", "Threads draft 2", "TikTok draft 2", "YouTube draft 2", 
            None
        ])
        # Row 4: another valid row with missing Link Post (None)
        ws1.append([
            3.0, 
            "Mobile payments trillion $ growth", 
            "./local_image.png", 
            "LinkedIn draft 3", "  ", " ", " ", " ", " ", " ", " ", 
            None
        ])
        
        # Sheet 2: Charity & Tokenization
        ws2 = wb.create_sheet(title="Charity & Tokenization")
        ws2.append([
            "#", "Trillion $ news Title", "Image link", "Linkedin", 
            "Facebook", "X (Twitter)", "Instagram", "Pinterest", 
            "Threads", "TikTok", "YouTube", "Link Post"
        ])
        ws2.append([
            1, "Charity Tokenization post", None, 
            "Draft for Linkedin", "Draft for Facebook", "Draft for X", 
            "Draft for Instagram", "Draft for Pinterest", "Draft for Threads", 
            "Draft for TikTok", "Draft for YouTube", None
        ])
        
        # Sheet 3: A helper/non-posting sheet (should be skipped based on missing signature header)
        ws3 = wb.create_sheet(title="Instructions")
        ws3.append(["Guide Title", "Description"])
        ws3.append(["How to run", "Step 1: search, Step 2: generate..."])
        
        wb.save(self.business_file)
        
        # Run ingestion
        rows, warnings = ingest_business_workbook(self.business_file)
        
        # Check rows parsed
        # Expecting:
        # - 2 rows from Payment (Row 2 and Row 4; Row 3 skipped because title was empty)
        # - 1 row from Charity & Tokenization (Row 2)
        # Total = 3 rows
        self.assertEqual(len(rows), 3)
        
        # Check row 1 details
        row1 = rows[0]
        self.assertEqual(row1.sheet_name, "Payment")
        self.assertEqual(row1.row_idx, 2)
        self.assertEqual(row1.id, 1)
        self.assertEqual(row1.title, "Merchant Payments trillion $ Opportunity")
        self.assertEqual(row1.image_link, "https://drive.google.com/file/123")
        self.assertEqual(row1.linkedin_draft, "LinkedIn draft content")
        self.assertEqual(row1.instagram_draft, None) # because it was "."
        self.assertEqual(row1.pinterest_draft, None) # because it was " "
        self.assertEqual(row1.tiktok_draft, None) # because it was "."
        self.assertEqual(row1.link_post_raw, "LinkedIn: https://linkedin.com/post1\nFacebook: https://facebook.com/post1")
        
        # Check row 2 details
        row2 = rows[1]
        self.assertEqual(row2.sheet_name, "Payment")
        self.assertEqual(row2.row_idx, 4)
        self.assertEqual(row2.id, 3)
        self.assertEqual(row2.title, "Mobile payments trillion $ growth")
        self.assertEqual(row2.facebook_draft, None) # because it was "  "
        self.assertEqual(row2.link_post_raw, None)
        
        # Check sheet 2 details
        row3 = rows[2]
        self.assertEqual(row3.sheet_name, "Charity & Tokenization")
        self.assertEqual(row3.row_idx, 2)
        self.assertEqual(row3.id, 1)
        self.assertEqual(row3.image_link, None)
        self.assertEqual(row3.linkedin_draft, "Draft for Linkedin")
        
        # Check warnings
        self.assertTrue(any("Row 3 in sheet 'Payment' skipped" in w for w in warnings))
        self.assertTrue(any("Instructions' does not appear to be a category posting sheet" in w for w in warnings))

    def test_ingest_business_workbook_missing_link_post_tolerated(self):
        # Create a business workbook where Link Post column is completely missing
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payment"
        ws.append([
            "#", "Trillion $ news Title", "Image link", "Linkedin", 
            "Facebook", "X (Twitter)", "Instagram", "Pinterest", 
            "Threads", "TikTok", "YouTube"
        ]) # No Link Post column
        ws.append([
            1, "Test Title", None, "Draft LI", None, None, None, None, None, None, None
        ])
        wb.save(self.business_file)
        
        rows, warnings = ingest_business_workbook(self.business_file)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].title, "Test Title")
        self.assertEqual(rows[0].linkedin_draft, "Draft LI")
        self.assertEqual(rows[0].link_post_raw, None)
        self.assertTrue(any("missing 'Link Post' column. It will be tolerated" in w for w in warnings))

    def test_ingest_business_workbook_missing_required_headers_raises_error(self):
        # Create a business workbook missing a required column (e.g. "Linkedin")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payment"
        ws.append([
            "#", "Trillion $ news Title", "Image link", 
            "Facebook", "X (Twitter)", "Instagram", "Pinterest", 
            "Threads", "TikTok", "YouTube", "Link Post"
        ]) # "Linkedin" is missing
        ws.append([
            1, "Test Title", None, "Draft FB", None, None, None, None, None, None, None
        ])
        wb.save(self.business_file)
        
        with self.assertRaises(ValueError) as context:
            ingest_business_workbook(self.business_file)
            
        self.assertIn("missing required category posting headers", str(context.exception))
        self.assertIn("linkedin", str(context.exception))

    def test_map_news_item_to_business_row(self):
        # Test mapping from legacy NewsItem to BusinessWorkbookRow
        item = NewsItem(
            id=42,
            title="Legacy News Title",
            image_file="image42.png",
            platform="linkedin",
            published_text="Social post draft text here",
            status="posted",
            notes="https://linkedin.com/post/42"
        )
        
        brow = map_news_item_to_business_row(item, row_idx=15)
        
        self.assertEqual(brow.sheet_name, "Compatibility")
        self.assertEqual(brow.row_idx, 15)
        self.assertEqual(brow.id, 42)
        self.assertEqual(brow.title, "Legacy News Title")
        self.assertEqual(brow.image_link, "image42.png")
        self.assertEqual(brow.linkedin_draft, "Social post draft text here")
        self.assertEqual(brow.facebook_draft, None)
        self.assertEqual(brow.link_post_raw, "LinkedIn: https://linkedin.com/post/42")

    def test_get_sheet_name_for_keyword(self):
        existing = ["Payment", "Charity & Tokenization", "Instructions"]
        self.assertEqual(get_sheet_name_for_keyword("Payment services", existing), "Payment")
        self.assertEqual(get_sheet_name_for_keyword("Charity tokenization", existing), "Charity & Tokenization")
        self.assertEqual(get_sheet_name_for_keyword("AI trillion dollar market", existing), "Ai Trillion Dollar Market")

    def test_save_news_to_business_excel(self):
        class MockConfig:
            def __init__(self, excel_file, image_dir):
                self.excel_file = excel_file
                self.image_dir = image_dir
                self.backup_enabled = False
                self.default_platform = "linkedin"

        # Setup config
        config = MockConfig(
            excel_file=os.path.join(self.test_dir, "test_save_news.xlsx"),
            image_dir=self.test_dir
        )
        
        # Create a temp screenshot
        temp_img_name = "temp_scr.png"
        with open(os.path.join(self.test_dir, temp_img_name), "w") as f:
            f.write("dummy screenshot content")
            
        items = [
            NewsItem(
                id=None,
                title="First Unique Title",
                keyword="Payment services",
                image_file=temp_img_name,
                url="https://test1.com"
            ),
            NewsItem(
                id=None,
                title="First Unique Title",  # Duplicate (should be skipped)
                keyword="Payment services",
                url="https://test1.com"
            ),
            NewsItem(
                id=None,
                title="Charity Tokenization Title",
                keyword="Charity & Tokenization",
                url="https://test2.com"
            )
        ]
        
        saved = save_news_to_business_excel(items, config)
        
        # Verify saved list size
        self.assertEqual(len(saved), 2)
        self.assertEqual(saved[0].title, "First Unique Title")
        self.assertEqual(saved[0].id, 1)
        self.assertEqual(saved[1].title, "Charity Tokenization Title")
        self.assertEqual(saved[1].id, 1)
        
        # Verify file exists
        self.assertTrue(os.path.exists(config.excel_file))
        
        # Load workbook and check sheet contents
        wb = openpyxl.load_workbook(config.excel_file)
        self.assertIn("Payment", wb.sheetnames)
        self.assertIn("Charity & Tokenization", wb.sheetnames)
        
        ws_pay = wb["Payment"]
        self.assertEqual(ws_pay.max_row, 2)  # Header + 1 row
        self.assertEqual(ws_pay.cell(row=2, column=1).value, 1)
        self.assertEqual(ws_pay.cell(row=2, column=2).value, "First Unique Title")
        
        # Check screenshot renamed
        from datetime import datetime
        today_str = datetime.now().strftime("%Y-%m-%d")
        expected_img_name = f"{today_str}_001_payment_services.png"
        self.assertEqual(ws_pay.cell(row=2, column=3).value, expected_img_name)
        self.assertTrue(os.path.exists(os.path.join(self.test_dir, expected_img_name)))
        
        # Check that the temp image file was deleted/moved
        self.assertFalse(os.path.exists(os.path.join(self.test_dir, temp_img_name)))
        
        wb.close()

    from unittest.mock import patch
    
    @patch("src.ai_writer.generate_ai_post")
    @patch("src.ai_writer.validate_generated_post")
    def test_generate_drafts_for_business_excel(self, mock_validate, mock_generate):
        class MockConfig:
            def __init__(self, excel_file, image_dir):
                self.excel_file = excel_file
                self.image_dir = image_dir
                self.backup_enabled = False
                self.default_platform = "linkedin"

        mock_generate.side_effect = lambda title, source, snippet, url, platform, config: f"Generated draft for {platform} - {title}"
        mock_validate.return_value = True
        
        # Create a business workbook with empty draft cells
        config = MockConfig(
            excel_file=os.path.join(self.test_dir, "test_generate_news.xlsx"),
            image_dir=self.test_dir
        )
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payment"
        headers = ["#", "Trillion $ news Title", "Image link", "Linkedin", "Facebook", "X (Twitter)", "Instagram", "Pinterest", "Threads", "TikTok", "YouTube", "Link Post"]
        ws.append(headers)
        # Add a row with missing LinkedIn draft
        ws.append([1, "Merchant Payments Title", None, None, "Existing Facebook draft", None, None, None, None, None, None, None])
        wb.save(config.excel_file)
        
        # Run generation
        generate_drafts_for_business_excel(config, limit=1, platform_option="linkedin")
        
        # Load and verify
        wb = openpyxl.load_workbook(config.excel_file)
        ws = wb["Payment"]
        # LinkedIn column is 4th (1-based index)
        self.assertEqual(ws.cell(row=2, column=4).value, "Generated draft for linkedin - Merchant Payments Title")
        # Facebook draft (5th column) remains untouched
        self.assertEqual(ws.cell(row=2, column=5).value, "Existing Facebook draft")
        wb.close()

if __name__ == "__main__":
    unittest.main()
