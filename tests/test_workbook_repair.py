import os
import tempfile
import unittest

import openpyxl

from src.business_workbook import ingest_business_workbook, repair_broken_draft_references


class TestWorkbookRepair(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.excel_file = os.path.join(self.test_dir, "Trillion $ news.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payment services"
        ws.append([
            "#",
            "Trillion $ news Title",
            "Image link",
            "Linkedin",
            "Facebook",
            "X (Twitter)",
            "Instagram",
            "Pinterest",
            "Threads",
            "TikTok",
            "YouTube",
            "Link Post",
        ])
        ws.append([
            6,
            "Vietnam banks prevent VND5 trillion in transfers after warnings",
            "2026-06-09_006_payment_services_trillion.png",
            None,
            "./output/posts/test_missing_workbook_repair_facebook.md",
            None,
            None,
            None,
            None,
            None,
            None,
            "Facebook: [posted-no-link]",
        ])
        wb.save(self.excel_file)
        wb.close()

        class MockConfig:
            backup_enabled = False
            excel_file = self.excel_file
            post_dir = os.path.join(self.test_dir, "posts")

        self.config = MockConfig()

    def test_ingest_marks_broken_draft_refs_without_exposing_fake_drafts(self):
        rows, warnings = ingest_business_workbook(self.excel_file)
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0].facebook_draft)
        self.assertEqual(
            rows[0].broken_draft_refs["facebook"],
            "./output/posts/test_missing_workbook_repair_facebook.md",
        )
        self.assertTrue(any("draft file not found" in warning for warning in warnings))

    def test_repair_clears_broken_draft_cells_in_excel(self):
        actions = repair_broken_draft_references(self.config)
        self.assertEqual(len(actions), 1)

        wb = openpyxl.load_workbook(self.excel_file)
        ws = wb["Payment services"]
        self.assertIsNone(ws.cell(row=2, column=5).value)
        wb.close()


if __name__ == "__main__":
    unittest.main()