import os
import shutil
import tempfile
import unittest
import openpyxl

from src.writeback import (
    is_workbook_locked,
    create_workbook_backup,
    evaluate_retry_disposition,
    write_post_result,
    WorkbookLockedError,
    BackupFailureError
)

class TestWriteback(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.workbook_path = os.path.join(self.test_dir, "Trillion $ news.xlsx")
        
        # Create a valid starter workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payment"
        ws.append([
            "#", "Trillion $ news Title", "Image link", "Linkedin", 
            "Facebook", "X (Twitter)", "Instagram", "Pinterest", 
            "Threads", "TikTok", "YouTube", "Link Post"
        ])
        # Row 2: has existing posting entries
        ws.append([
            1, "Title 1", None, "Draft LI", None, None, None, None, None, None, None,
            "LinkedIn: https://linkedin.com/1\nFacebook: [error] Auth failed"
        ])
        # Row 3: has no Link Post column value (None)
        ws.append([
            2, "Title 2", None, "Draft LI", None, None, None, None, None, None, None,
            None
        ])
        
        # Sheet 2: Charity
        ws2 = wb.create_sheet("Charity")
        ws2.append([
            "#", "Trillion $ news Title", "Image link", "Linkedin", 
            "Facebook", "X (Twitter)", "Instagram", "Pinterest", 
            "Threads", "TikTok", "YouTube"
        ]) # Lacks "Link Post" header
        ws2.append([
            1, "Charity Title", None, "Draft LI", None, None, None, None, None, None, None
        ])
        
        wb.save(self.workbook_path)
        
    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def test_is_workbook_locked(self):
        self.assertFalse(is_workbook_locked(self.workbook_path))
        self.assertFalse(is_workbook_locked("nonexistent_file.xlsx"))
        
        dir_name = os.path.dirname(self.workbook_path) or "."
        base_name = os.path.basename(self.workbook_path)
        
        # 1. Test LibreOffice lock file
        libreoffice_lock = os.path.join(dir_name, f".~lock.{base_name}#")
        with open(libreoffice_lock, "w") as f:
            f.write("lock details")
        try:
            self.assertTrue(is_workbook_locked(self.workbook_path))
        finally:
            if os.path.exists(libreoffice_lock):
                os.remove(libreoffice_lock)
                
        # 2. Test MS Office lock file
        ms_office_lock = os.path.join(dir_name, f"~${base_name}")
        with open(ms_office_lock, "w") as f:
            f.write("lock details")
        try:
            self.assertTrue(is_workbook_locked(self.workbook_path))
        finally:
            if os.path.exists(ms_office_lock):
                os.remove(ms_office_lock)

    def test_resolve_workbook_path_rejects_backup_filename(self):
        from src.writeback import resolve_workbook_path, WORKBOOK_BACKUP_DIRNAME

        canonical = resolve_workbook_path(self.test_dir)
        backup_like = os.path.join(self.test_dir, "Trillion $ news.20260609_090724.xlsx")
        self.assertEqual(resolve_workbook_path(self.test_dir, backup_like), canonical)

    def test_create_workbook_backup(self):
        from src.writeback import WORKBOOK_BACKUP_DIRNAME

        backup_path = create_workbook_backup(self.workbook_path)
        self.assertTrue(os.path.exists(backup_path))
        self.assertIn(WORKBOOK_BACKUP_DIRNAME, backup_path)
        self.assertIn("Trillion $ news.", backup_path)
        self.assertTrue(backup_path.endswith(".xlsx"))
        
        # Ensure copy is identical
        wb_orig = openpyxl.load_workbook(self.workbook_path)
        wb_backup = openpyxl.load_workbook(backup_path)
        self.assertEqual(wb_orig.sheetnames, wb_backup.sheetnames)
        wb_orig.close()
        wb_backup.close()

    def test_evaluate_retry_disposition(self):
        # success/skip -> False
        self.assertFalse(evaluate_retry_disposition("LinkedIn: https://linkedin.com/1", "linkedin"))
        self.assertFalse(evaluate_retry_disposition("X: [posted-no-link]", "x"))
        self.assertFalse(evaluate_retry_disposition("Facebook: [skip] duplicate", "facebook"))
        
        # retryable/missing -> True
        self.assertTrue(evaluate_retry_disposition("LinkedIn: [error] Auth failed", "linkedin"))
        self.assertTrue(evaluate_retry_disposition("LinkedIn: [pending]", "linkedin"))
        self.assertTrue(evaluate_retry_disposition("LinkedIn: [login-required]", "linkedin"))
        self.assertTrue(evaluate_retry_disposition(None, "linkedin"))
        self.assertTrue(evaluate_retry_disposition("", "linkedin"))
        self.assertTrue(evaluate_retry_disposition("LinkedIn: https://linkedin.com/1", "facebook")) # Facebook is missing, so retryable

    def test_write_post_result_success(self):
        # 1. Update Facebook (retryable -> success URL) on Row 2
        # Verify LinkedIn is preserved
        new_val = write_post_result(
            workbook_path=self.workbook_path,
            sheet_name="Payment",
            row_idx=2,
            platform="facebook",
            status_value="https://facebook.com/post/123",
            backup_enabled=False
        )
        
        expected = "Facebook: https://facebook.com/post/123\nLinkedIn: https://linkedin.com/1"
        self.assertEqual(new_val, expected)
        
        # Verify in sheet
        wb = openpyxl.load_workbook(self.workbook_path)
        ws = wb["Payment"]
        self.assertEqual(ws.cell(row=2, column=12).value, expected)
        wb.close()

    def test_write_post_result_creates_column(self):
        # Sheet 'Charity' has no 'Link Post' column initially
        # Updating it should create the column in column 12
        new_val = write_post_result(
            workbook_path=self.workbook_path,
            sheet_name="Charity",
            row_idx=2,
            platform="linkedin",
            status_value="https://linkedin.com/charity/1",
            backup_enabled=False
        )
        
        expected = "LinkedIn: https://linkedin.com/charity/1"
        self.assertEqual(new_val, expected)
        
        wb = openpyxl.load_workbook(self.workbook_path)
        ws = wb["Charity"]
        # Header updated
        self.assertEqual(ws.cell(row=1, column=12).value, "Link Post")
        # Cell updated
        self.assertEqual(ws.cell(row=2, column=12).value, expected)
        wb.close()

    def test_write_post_result_backups(self):
        # Trigger with backup enabled
        write_post_result(
            workbook_path=self.workbook_path,
            sheet_name="Payment",
            row_idx=3,
            platform="x",
            status_value="[posted-no-link]",
            backup_enabled=True
        )
        
        from src.writeback import WORKBOOK_BACKUP_DIRNAME, is_timestamped_workbook_backup

        backup_dir = os.path.join(self.test_dir, WORKBOOK_BACKUP_DIRNAME)
        self.assertTrue(os.path.isdir(backup_dir))
        backups = [
            f for f in os.listdir(backup_dir)
            if is_timestamped_workbook_backup(f)
        ]
        self.assertEqual(len(backups), 1)

    def test_write_post_result_failures(self):
        # Missing sheet
        with self.assertRaises(ValueError) as context:
            write_post_result(self.workbook_path, "NonexistentSheet", 2, "linkedin", "https://link", False)
        self.assertIn("Sheet 'NonexistentSheet' not found", str(context.exception))
        
        # Row out of bounds
        with self.assertRaises(ValueError) as context2:
            write_post_result(self.workbook_path, "Payment", 50, "linkedin", "https://link", False)
        self.assertIn("Row index 50 is out of bounds", str(context2.exception))
        
        # Invalid status value format
        with self.assertRaises(ValueError) as context3:
            write_post_result(self.workbook_path, "Payment", 2, "linkedin", "plain text not allowed", False)
        self.assertIn("Invalid status value", str(context3.exception))

if __name__ == "__main__":
    unittest.main()
