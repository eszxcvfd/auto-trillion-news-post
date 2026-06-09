import os
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

from src.business_workbook import save_news_to_business_excel, generate_drafts_for_business_excel
from src.models import NewsItem


class TestDraftRunWorkbookBackfill(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.excel_file = os.path.join(self.test_dir, "Trillion $ news.xlsx")
        self.image_dir = os.path.join(self.test_dir, "images")
        os.makedirs(self.image_dir, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payment services"
        ws.append([
            "#",
            "Trillion $ news Title",
            "Image link",
            "Linkedin",
            "Facebook",
            "X (Twitter)",
            "Instagram",
            "Pinterest",
            "Threads",
            "TikTok",
            "YouTube",
            "Link Post",
        ])
        ws.append([
            1,
            "Existing trillion $ article",
            None,
            "./output/posts/missing_existing_linkedin.md",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ])
        wb.save(self.excel_file)
        wb.close()

        class MockConfig:
            backup_enabled = False
            default_platform = "linkedin"

            def __init__(self, excel_file, image_dir):
                self.excel_file = excel_file
                self.image_dir = image_dir

        self.config = MockConfig(self.excel_file, self.image_dir)

    def test_duplicate_harvest_backfills_missing_image_link(self):
        temp_image = os.path.join(self.image_dir, "temp_capture.png")
        with open(temp_image, "wb") as handle:
            handle.write(b"png")

        item = NewsItem(
            id=99,
            title="Existing trillion $ article",
            keyword="Payment services",
            image_file="temp_capture.png",
            found_date="2026-06-09",
        )

        saved = save_news_to_business_excel([item], self.config, limit=2)
        self.assertEqual(saved, [])

        wb = openpyxl.load_workbook(self.excel_file)
        ws = wb["Payment services"]
        image_val = ws.cell(row=2, column=3).value
        self.assertEqual(image_val, "2026-06-09_001_payment_services.png")
        self.assertTrue(os.path.exists(os.path.join(self.image_dir, image_val)))
        wb.close()

    @patch("src.ai_writer.validate_generated_post", return_value=True)
    @patch("src.ai_writer.generate_ai_post")
    def test_generate_all_missing_platform_drafts(self, mock_generate, _mock_validate):
        mock_generate.side_effect = (
            lambda title, source, snippet, url, platform, config: f"Draft for {platform}"
        )

        generate_drafts_for_business_excel(self.config, limit=None, platform_option="all")

        wb = openpyxl.load_workbook(self.excel_file)
        ws = wb["Payment services"]
        from src.draft_paths import resolve_post_draft_path

        linkedin_cell = str(ws.cell(row=2, column=4).value)
        self.assertTrue(linkedin_cell.endswith("_linkedin.md"))
        self.assertTrue(os.path.exists(resolve_post_draft_path(linkedin_cell, self.excel_file)))
        self.assertIsNone(ws.cell(row=2, column=5).value)
        self.assertIsNone(ws.cell(row=2, column=6).value)
        mock_generate.assert_called_once()
        wb.close()


if __name__ == "__main__":
    unittest.main()
