import os
import shutil
import tempfile
import unittest

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    openpyxl = None
    OPENPYXL_AVAILABLE = False

from src.models import NewsItem
from src.config import AppConfig
from src.excel_store import (
    slugify,
    init_excel_file,
    get_next_id,
    save_news_to_excel,
    HEADERS
)

@unittest.skipIf(not OPENPYXL_AVAILABLE, "openpyxl is not installed")
class TestExcelStore(unittest.TestCase):
    def setUp(self):
        # Create temp folder for test output
        self.test_dir = tempfile.mkdtemp()
        self.config = AppConfig()
        self.config.excel_file = os.path.join(self.test_dir, "Trillion $ test news.xlsx")
        self.config.image_dir = os.path.join(self.test_dir, "images")
        os.makedirs(self.config.image_dir, exist_ok=True)

    def tearDown(self):
        shutil.rmtree(self.test_dir)

    def test_slugify(self):
        self.assertEqual(slugify("Payment services trillion $"), "payment_services_trillion")
        self.assertEqual(slugify("AI trillion-dollar market!!!"), "ai_trillion_dollar_market")

    def test_init_excel_file(self):
        init_excel_file(self.config.excel_file)
        self.assertTrue(os.path.exists(self.config.excel_file))
        
        # Verify headers
        wb = openpyxl.load_workbook(self.config.excel_file)
        ws = wb.active
        self.assertEqual(ws.max_row, 1)
        for i, h in enumerate(HEADERS):
            self.assertEqual(ws.cell(row=1, column=i+1).value, h)

    def test_get_next_id(self):
        init_excel_file(self.config.excel_file)
        wb = openpyxl.load_workbook(self.config.excel_file)
        ws = wb.active
        
        self.assertEqual(get_next_id(ws), 1)
        
        ws.append([1, "2026-06-04", "kw", "title", "src", "url"])
        self.assertEqual(get_next_id(ws), 2)
        
        ws.append([2, "2026-06-04", "kw", "title", "src", "url"])
        self.assertEqual(get_next_id(ws), 3)

    def test_save_news_to_excel_appends_and_renames(self):
        # 1. Create a dummy temp image to simulate search capture
        temp_img_name = "temp_123456.png"
        temp_img_path = os.path.join(self.config.image_dir, temp_img_name)
        with open(temp_img_path, "w") as f:
            f.write("dummy image data")
            
        items = [
            NewsItem(
                keyword="AI trillion",
                title="AI trillion dollar market",
                url="https://example.com/ai",
                image_file=temp_img_name
            )
        ]
        
        # Save
        saved = save_news_to_excel(items, self.config)
        self.assertEqual(len(saved), 1)
        
        item = saved[0]
        self.assertEqual(item.id, 1)
        self.assertEqual(item.status, "new")
        
        # Verify image was renamed
        expected_img_name = f"{item.found_date}_001_ai_trillion.png"
        self.assertEqual(item.image_file, expected_img_name)
        self.assertTrue(os.path.exists(os.path.join(self.config.image_dir, expected_img_name)))
        self.assertFalse(os.path.exists(temp_img_path)) # Old temp should be gone

        # 2. Try to save a duplicate URL
        duplicate_items = [
            NewsItem(
                keyword="AI trillion",
                title="AI trillion dollar market",
                url="https://example.com/ai"
            )
        ]
        saved_dup = save_news_to_excel(duplicate_items, self.config)
        self.assertEqual(len(saved_dup), 0) # Should be skipped since URL already exists in excel

if __name__ == "__main__":
    unittest.main()
