import os
import openpyxl
import subprocess

excel_path = "output/Trillion $ news.xlsx"
if not os.path.exists(excel_path):
    print("[ERROR] Excel file not found. Please run search/generate first.")
    exit(1)

wb = openpyxl.load_workbook(excel_path)
ws = wb.active

generated_ids = []
for r in range(2, ws.max_row + 1):
    status = ws.cell(row=r, column=13).value
    id_val = ws.cell(row=r, column=1).value
    if status == "generated":
        generated_ids.append(id_val)

wb.close()

if not generated_ids:
    print("[INFO] No articles with status 'generated' found in the Excel file.")
    exit(0)

print(f"[INFO] Found {len(generated_ids)} articles ready to post: {generated_ids}")
for item_id in generated_ids:
    print(f"\n==========================================")
    print(f"   STARTING POST FOR ARTICLE ID: {item_id}")
    print(f"==========================================\n")
    
    # Run the main.py post command for this ID
    # Use python3 (will automatically run in activated virtualenv if active)
    subprocess.run(["python3", "main.py", "post", "--id", str(item_id)], check=False)
