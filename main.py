import os
import argparse
import sys
import json
from src.config import AppConfig
from src.models import NewsItem
from src.searcher import search_news, sync_playwright as search_sync_playwright
from src.filter import is_trillion_news, deduplicate_news

# Embedded defaults
DEFAULT_ENV_EXAMPLE = """# Gemini API Configuration
GEMINI_API_KEY=your_gemini_api_key_here
AI_PROVIDER=gemini
AI_MODEL=gemini-1.5-flash

# Search Configuration
SEARCH_PROVIDER=bing
HEADLESS=false

# Output Configuration
OUTPUT_DIR=./output
EXCEL_FILE=./output/Trillion $ news.xlsx
IMAGE_DIR=./output/Ảnh Trillion $ news
POST_DIR=./output/posts
LOG_DIR=./output/logs

# Run Limits & Preferences
DEFAULT_PLATFORM=linkedin
DEFAULT_LANGUAGE=en
MAX_RESULTS_PER_KEYWORD=10
MAX_POSTS_PER_RUN=5
"""

DEFAULT_CONFIG_YAML = """# Default System Configuration
platform: linkedin
language: en

search:
  provider: bing
  max_results_per_keyword: 10
  require_terms:
    - trillion
    - trillion-dollar
    - "$ trillion"
    - "USD"
  exclude_domains:
    - reddit.com
    - x.com
    - facebook.com

hashtags:
  fixed_bottom:
    - TAHKFoundation
    - HenryUniverses
    - USIran
    - USTariffs
    - Trump

posting:
  mode: draft_only
  human_confirm_before_post: true
"""

DEFAULT_KEYWORDS_TXT = """Payment services trillion $
Mobile payments trillion $
Fintech trillion $
AI trillion dollar market
Healthcare trillion dollar market
Energy trillion dollar market
Real estate trillion dollar market
Banking trillion dollar market
Cross border payments trillion $
Digital payments trillion $
"""

def init_project():
    print("Initializing project scaffolding...")
    
    # Files to create if they don't exist
    files_to_create = {
        ".env.example": DEFAULT_ENV_EXAMPLE,
        "config.yaml": DEFAULT_CONFIG_YAML,
        "keywords.txt": DEFAULT_KEYWORDS_TXT
    }
    
    for filename, content in files_to_create.items():
        if os.path.exists(filename):
            print(f"[INFO] {filename} already exists, skipping.")
        else:
            with open(filename, "w", encoding="utf-8") as f:
                f.write(content)
            print(f"[SUCCESS] Created {filename}")
            
    # Directories to create
    directories = [
        "output",
        "output/posts",
        "output/Ảnh Trillion $ news",
        "output/logs"
    ]
    
    for directory in directories:
        if os.path.exists(directory):
            print(f"[INFO] Directory {directory} already exists, skipping.")
        else:
            os.makedirs(directory, exist_ok=True)
            print(f"[SUCCESS] Created directory {directory}")
            
    # Initialize Business Workbook if it doesn't exist
    config = AppConfig()
    filepath = config.excel_file
    if os.path.exists(filepath):
        print(f"[INFO] Workbook '{filepath}' already exists, skipping.")
    else:
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Payment"
            headers = ["#", "Trillion $ news Title", "Image link", "Linkedin", "Facebook", "X (Twitter)", "Instagram", "Pinterest", "Threads", "TikTok", "YouTube", "Link Post"]
            ws.append(headers)
            # Create a secondary default category sheet
            ws2 = wb.create_sheet(title="Charity & Tokenization")
            ws2.append(headers)
            wb.save(filepath)
            print(f"[SUCCESS] Initialized new Business Workbook at {filepath}")
        except Exception as e:
            print(f"[WARNING] Failed to initialize Excel workbook: {e}")
            
    print("Project initialization complete.")

def load_keywords(filepath: str) -> list[str]:
    if not os.path.exists(filepath):
        print(f"[ERROR] Keywords file '{filepath}' does not exist.")
        return []
    keywords = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            keywords.append(line)
    return keywords

def execute_search(keywords_path: str, limit: int = None):
    if search_sync_playwright is None:
        print("[ERROR] Playwright is not installed. Please run pip install -r requirements.txt to install it.")
        return

    # Load configuration
    config = AppConfig()
    
    # Pre-clean existing temp files in the image directory from previous runs
    if os.path.exists(config.image_dir):
        for f in os.listdir(config.image_dir):
            if f.startswith("temp_"):
                try:
                    os.remove(os.path.join(config.image_dir, f))
                except Exception:
                    pass
    
    # Load keywords
    keywords = load_keywords(keywords_path)
    if not keywords:
        print("[ERROR] No keywords found to search.")
        sys.exit(1)
        
    print(f"[INFO] Loaded {len(keywords)} keywords.")
    
    all_raw_news = []
    for kw in keywords:
        raw_news = search_news(kw, config)
        all_raw_news.extend(raw_news)
        
    # Filter
    filtered_news = [item for item in all_raw_news if is_trillion_news(item, config)]
    print(f"[INFO] {len(filtered_news)} / {len(all_raw_news)} articles passed trillion filter.")
    
    # Deduplicate
    unique_news = deduplicate_news(filtered_news)
    print(f"[INFO] {len(unique_news)} articles remaining after deduplication.")
    
    # Determine the save limit (to avoid saving too many new rows)
    save_limit = limit
    if save_limit is None and config.max_posts_per_run is not None:
        save_limit = config.max_posts_per_run
        
    # Clean up unsaved temp files
    retained_temp_files = {item.image_file for item in unique_news if item.image_file and item.image_file.startswith("temp_")}
    for item in all_raw_news:
        if item.image_file and item.image_file.startswith("temp_"):
            if item.image_file not in retained_temp_files:
                temp_path = os.path.join(config.image_dir, item.image_file)
                if os.path.exists(temp_path):
                    try:
                        os.remove(temp_path)
                    except Exception:
                        pass
        
    # Save to Excel and rename temp screenshots
    filepath = config.excel_file
    wb_type = "business"  # default if file doesn't exist
    if os.path.exists(filepath):
        from src.business_workbook import detect_workbook_type
        try:
            wb_type = detect_workbook_type(filepath)
        except Exception:
            wb_type = "legacy"
            
    if wb_type == "business":
        from src.business_workbook import save_news_to_business_excel
        saved_news = save_news_to_business_excel(unique_news, config, limit=save_limit)
    else:
        from src.excel_store import save_news_to_excel
        saved_news = save_news_to_excel(unique_news, config, limit=save_limit)
        
    # Final foolproof cleanup of all temp files that were not saved
    saved_filenames = {item.image_file for item in saved_news if item.image_file}
    if os.path.exists(config.image_dir):
        for f in os.listdir(config.image_dir):
            if f.startswith("temp_") and f not in saved_filenames:
                try:
                    os.remove(os.path.join(config.image_dir, f))
                except Exception:
                    pass
    
    # Format and print JSON
    dict_news = [item.to_dict() for item in saved_news]
    print(json.dumps(dict_news, indent=2, ensure_ascii=False))
    return saved_news

def normalize_generate_platform(platform: str = None) -> str:
    """Blank platform means generate drafts for every supported platform."""
    if platform is None:
        return "all"
    cleaned = str(platform).strip().lower()
    if not cleaned or cleaned == "none":
        return "all"
    return platform


def execute_generate(platform: str = None, limit: int = None, target_ids: list = None):
    config = AppConfig()
    
    # Load openpyxl safely
    try:
        import openpyxl
    except ImportError:
        print("[ERROR] openpyxl is not installed. Please run pip install -r requirements.txt to install it.")
        sys.exit(1)
        
    platform = normalize_generate_platform(platform)
    
    filepath = config.excel_file
    if not os.path.exists(filepath):
        print(f"[ERROR] Excel file '{filepath}' does not exist. Please run search first.")
        sys.exit(1)
        
    from src.business_workbook import detect_workbook_type
    try:
        wb_type = detect_workbook_type(filepath)
    except Exception:
        wb_type = "legacy"
        
    if wb_type == "business":
        from src.business_workbook import generate_drafts_for_business_excel
        generate_drafts_for_business_excel(config, limit=limit, platform_option=platform, target_ids=target_ids)
        return
        
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
    except Exception as e:
        print(f"[ERROR] Failed to load Excel workbook: {e}")
        sys.exit(1)
        
    # Read rows with status "new"
    new_items = []
    for r in range(2, ws.max_row + 1):
        id_val = ws.cell(row=r, column=1).value
        try:
            current_id = int(id_val)
        except (ValueError, TypeError):
            current_id = id_val
            
        if target_ids is not None and current_id not in target_ids:
            continue
            
        status_val = ws.cell(row=r, column=13).value
        if status_val == "new":
            item = NewsItem(
                id=current_id,
                found_date=ws.cell(row=r, column=2).value,
                keyword=ws.cell(row=r, column=3).value,
                title=ws.cell(row=r, column=4).value,
                source=ws.cell(row=r, column=5).value,
                url=ws.cell(row=r, column=6).value,
                snippet=ws.cell(row=r, column=7).value,
                published_text=ws.cell(row=r, column=8).value,
                image_file=ws.cell(row=r, column=9).value,
                platform=platform,
                status="new",
                notes=ws.cell(row=r, column=14).value
            )
            new_items.append(item)
            
    wb.close()
    
    if not new_items:
        print("[INFO] No new items found in Excel to generate posts for.")
        return
        
    print(f"[INFO] Found {len(new_items)} new articles to process.")
    
    # Apply limit per keyword
    if limit is not None:
        limited_items = []
        kw_counts = {}
        for item in new_items:
            kw = item.keyword or "Payment"
            count = kw_counts.get(kw, 0)
            if count < limit:
                limited_items.append(item)
                kw_counts[kw] = count + 1
        new_items = limited_items
    elif config.max_posts_per_run is not None:
        new_items = new_items[:config.max_posts_per_run]
        
    print(f"[INFO] Processing {len(new_items)} articles...")
    
    from src.ai_writer import generate_ai_post, validate_generated_post
    from src.post_writer import write_post_file, update_excel_row_with_post, extract_top_hashtags
    
    for item in new_items:
        print(f"[INFO] Generating post for ID {item.id}: '{item.title}'")
        try:
            content = generate_ai_post(
                title=item.title,
                source=item.source or "",
                snippet=item.snippet or "",
                url=item.url or "",
                platform=platform,
                config=config
            )
            
            if not validate_generated_post(content, platform):
                print(f"[WARNING] Post format validation failed for ID {item.id}.")
                update_excel_row_with_post(
                    item_id=item.id,
                    post_file_path="",
                    top_hashtags="",
                    status="error",
                    config=config,
                    notes="Validation failed: post format did not match requirements."
                )
                continue
                
            post_file = write_post_file(item, content, config)
            top_tags = extract_top_hashtags(content)
            
            update_excel_row_with_post(
                item_id=item.id,
                post_file_path=post_file,
                top_hashtags=top_tags,
                status="generated",
                config=config
            )
            
        except Exception as e:
            print(f"[ERROR] Failed to generate post for ID {item.id}: {e}")
            update_excel_row_with_post(
                item_id=item.id,
                post_file_path="",
                top_hashtags="",
                status="error",
                config=config,
                notes=f"Generation failed: {e}"
            )

def execute_run(keywords_path: str, platform: str = None, limit: int = None):
    missing_dependencies = []

    if search_sync_playwright is None:
        missing_dependencies.append("Playwright")

    try:
        import openpyxl
    except ImportError:
        missing_dependencies.append("openpyxl")

    if missing_dependencies:
        for dependency in missing_dependencies:
            print(f"[ERROR] {dependency} is not installed. Please run pip install -r requirements.txt to install it.")
        return

    print("=== Starting Full Draft Pipeline ===")
    saved_news = execute_search(keywords_path, limit)
    
    # Check workbook type to format target_ids correctly
    config = AppConfig()
    filepath = config.excel_file
    wb_type = "business"
    if os.path.exists(filepath):
        from src.business_workbook import detect_workbook_type
        try:
            wb_type = detect_workbook_type(filepath)
        except Exception:
            wb_type = "legacy"
            
    target_ids = None
    if saved_news:
        if wb_type == "business":
            target_ids = [(item.keyword, item.id) for item in saved_news]
        else:
            target_ids = [item.id for item in saved_news]

    # Harvest limit applies only to newly saved rows. When nothing new was saved,
    # backfill every row that still has missing platform drafts.
    generate_limit = limit if target_ids else None
    execute_generate(
        normalize_generate_platform(platform),
        generate_limit,
        target_ids=target_ids,
    )
    print("=== Full Draft Pipeline Completed ===")

def execute_post(item_id: int, platform: str = None, sheet: str = None):
    config = AppConfig()
    
    try:
        import openpyxl
    except ImportError:
        print("[ERROR] openpyxl is not installed. Please run pip install -r requirements.txt to install it.")
        sys.exit(1)
    
    filepath = config.excel_file
    if not os.path.exists(filepath):
        print(f"[ERROR] Excel file '{filepath}' does not exist. Please run search/generate first.")
        sys.exit(1)
        
    from src.business_workbook import detect_workbook_type
    try:
        wb_type = detect_workbook_type(filepath)
    except Exception:
        wb_type = "legacy"
        
    if wb_type == "business":
        from src.business_workbook import ingest_business_workbook
        from src.posting_core import (
            canonicalize_platform,
            list_eligible_platforms,
            resolve_manual_post_target,
            PLATFORM_CANONICAL,
        )
        from src.writeback import write_post_result
        from src.assisted_posting import run_assisted_posting
        
        sheet_scope = [sheet] if sheet else None
        try:
            rows, warnings = ingest_business_workbook(filepath, sheet_scope=sheet_scope)
        except Exception as e:
            print(f"[ERROR] Failed to ingest business workbook: {e}")
            sys.exit(1)
            
        matching_rows = [r for r in rows if r.id == item_id]
        if not matching_rows:
            print(f"[ERROR] Row with ID {item_id} not found in the Business Workbook.")
            sys.exit(1)
            
        row = matching_rows[0]
        if len(matching_rows) > 1:
            print(f"[INFO] Multiple rows matching ID {item_id} found across sheets:")
            for idx, r in enumerate(matching_rows):
                print(f"  {idx + 1}. Sheet: '{r.sheet_name}', Title: '{r.title}'")
            try:
                choice = input(f"Select which row you want to post (1-{len(matching_rows)}, default 1): ").strip()
                if choice:
                    choice_idx = int(choice) - 1
                    if 0 <= choice_idx < len(matching_rows):
                        row = matching_rows[choice_idx]
                        print(f"[INFO] Selected match in sheet '{row.sheet_name}'.")
                    else:
                        print(f"[WARNING] Invalid selection. Using default match in sheet '{row.sheet_name}'.")
                else:
                    print(f"[INFO] No selection. Using default match in sheet '{row.sheet_name}'.")
            except (EOFError, ValueError, IndexError, KeyboardInterrupt):
                print(f"[INFO] Using default match in sheet '{row.sheet_name}'.")
            
        eligible = list_eligible_platforms(row, config=config, workbook_path=filepath)
        if not eligible:
            print(
                f"[ERROR] No eligible platforms to post for row ID {item_id} "
                f"in sheet '{row.sheet_name}'. Ensure draft content exists and "
                "the platform is not already posted or explicitly skipped."
            )
            sys.exit(1)

        platform_key = None
        draft_content = None
        if platform:
            try:
                platform_key, draft_content = resolve_manual_post_target(
                    row,
                    platform,
                    config=config,
                    workbook_path=filepath,
                )
            except ValueError as e:
                print(f"[ERROR] {e}")
                sys.exit(1)
        elif len(eligible) == 1:
            platform_key, draft_content = eligible[0]
        else:
            print(f"[INFO] Multiple eligible platforms found for row ID {item_id}:")
            for idx, (p_key, _) in enumerate(eligible, 1):
                print(f"  {idx}. {PLATFORM_CANONICAL[p_key]}")
            try:
                choice = input(
                    f"Select platform to post (1-{len(eligible)}, default 1): "
                ).strip()
                choice_idx = int(choice) - 1 if choice else 0
                if not (0 <= choice_idx < len(eligible)):
                    print("[ERROR] Invalid platform selection.")
                    sys.exit(1)
                platform_key, draft_content = eligible[choice_idx]
                print(f"[INFO] Selected platform '{PLATFORM_CANONICAL[platform_key]}'.")
            except (EOFError, ValueError, KeyboardInterrupt):
                print("[ERROR] Platform selection is required when multiple eligible platforms exist.")
                sys.exit(1)

        platform = platform_key
            
        # Resolve draft_content if it is stored as a file path
        if draft_content and isinstance(draft_content, str):
            test_paths = [
                draft_content,
                os.path.join(os.path.dirname(filepath), draft_content) if filepath else "",
                os.path.join(config.post_dir, os.path.basename(draft_content)) if hasattr(config, "post_dir") else ""
            ]
            for p in test_paths:
                if p and os.path.exists(p) and os.path.isfile(p):
                    from src.assisted_posting import parse_post_markdown
                    parsed = parse_post_markdown(p)
                    if parsed:
                        draft_content = parsed
                        break
            
        # Construct NewsItem wrapper for run_assisted_posting
        item = NewsItem(
            id=row.id,
            title=row.title,
            image_file=row.image_link,
            platform=platform,
            status="new"
        )
        
        print(f"[INFO] Launching assisted posting for ID {item_id} on '{platform}' [Sheet: {row.sheet_name}, Row: {row.row_idx}]: '{item.title}'")
        success = run_assisted_posting(
            item,
            config,
            post_content=draft_content,
            workbook_path=filepath,
        )
        
        if success:
            url_input = getattr(item, "post_url", "").strip()
            status_val = url_input if url_input else "[posted-no-link]"
            
            try:
                write_post_result(
                    workbook_path=filepath,
                    sheet_name=row.sheet_name,
                    row_idx=row.row_idx,
                    platform=platform,
                    status_value=status_val,
                    backup_enabled=config.backup_enabled
                )
                print(f"[SUCCESS] Post result successfully written to Link Post in sheet '{row.sheet_name}', row {row.row_idx}.")
            except Exception as e:
                print(f"[ERROR] Failed to write post result back to Excel: {e}")
        else:
            print(f"[INFO] Post {item_id} was not marked as posted.")
            ans_skip = input("Do you want to mark this item as skipped? [y/N]: ").strip().lower()
            if ans_skip in ["y", "yes"]:
                try:
                    write_post_result(
                        workbook_path=filepath,
                        sheet_name=row.sheet_name,
                        row_idx=row.row_idx,
                        platform=platform,
                        status_value="[skip] skipped by operator",
                        backup_enabled=config.backup_enabled
                    )
                    print(f"[SUCCESS] Post successfully marked as skipped in sheet '{row.sheet_name}', row {row.row_idx}.")
                except Exception as e:
                    print(f"[ERROR] Failed to write skip status to Excel: {e}")
                    
    else:
        # Legacy Contract A flow
        platform = platform or config.default_platform
        try:
            wb = openpyxl.load_workbook(filepath)
            ws = wb.active
        except Exception as e:
            print(f"[ERROR] Failed to load Excel workbook: {e}")
            sys.exit(1)
            
        item = None
        for r in range(2, ws.max_row + 1):
            id_val = ws.cell(row=r, column=1).value
            try:
                current_id = int(id_val)
            except (ValueError, TypeError):
                current_id = id_val
                
            if current_id == item_id:
                item = NewsItem(
                    id=current_id,
                    found_date=ws.cell(row=r, column=2).value,
                    keyword=ws.cell(row=r, column=3).value,
                    title=ws.cell(row=r, column=4).value,
                    source=ws.cell(row=r, column=5).value,
                    url=ws.cell(row=r, column=6).value,
                    snippet=ws.cell(row=r, column=7).value,
                    published_text=ws.cell(row=r, column=8).value,
                    image_file=ws.cell(row=r, column=9).value,
                    platform=platform,
                    top_hashtags=ws.cell(row=r, column=11).value,
                    generated_post_file=ws.cell(row=r, column=12).value,
                    status=ws.cell(row=r, column=13).value,
                    notes=ws.cell(row=r, column=14).value
                )
                break
                
        wb.close()
        
        if not item:
            print(f"[ERROR] Article with ID {item_id} not found in Excel.")
            sys.exit(1)
            
        if not item.generated_post_file or not os.path.exists(item.generated_post_file):
            print(f"[ERROR] No generated post file found for ID {item_id}. Please run generate first.")
            sys.exit(1)
            
        from src.assisted_posting import run_assisted_posting
        from src.post_writer import update_excel_row_with_post
        
        print(f"[INFO] Launching assisted posting for ID {item_id}: '{item.title}'")
        success = run_assisted_posting(item, config)
        
        if success:
            update_excel_row_with_post(item_id=item_id, status="posted", config=config)
            print(f"[SUCCESS] Post {item_id} successfully marked as 'posted' in Excel.")
        else:
            print(f"[INFO] Post {item_id} was not marked as posted.")
            ans_skip = input("Do you want to mark this item as skipped? [y/N]: ").strip().lower()
            if ans_skip in ["y", "yes"]:
                update_excel_row_with_post(item_id=item_id, status="skipped", config=config)
                print(f"[SUCCESS] Post {item_id} successfully marked as 'skipped' in Excel.")

def execute_inspect_workbook(
    workbook_path: str = None, 
    sheet: str = None, 
    as_json: bool = False,
    as_plan: bool = False,
    limit: int = 2,
    platforms_str: str = ""
):
    from src.business_workbook import detect_workbook_type, ingest_business_workbook
    from src.posting_core import build_posting_plan
    config = AppConfig()
    filepath = workbook_path or config.excel_file
    
    if not os.path.exists(filepath):
        print(f"[ERROR] Excel file '{filepath}' does not exist.")
        sys.exit(1)
        
    try:
        wb_type = detect_workbook_type(filepath)
    except Exception as e:
        print(f"[ERROR] Detection failed: {e}")
        sys.exit(1)
        
    print(f"Workbook Path: {filepath}")
    print(f"Detected Type: {wb_type.upper()} WORKBOOK (Contract {'A' if wb_type == 'legacy' else 'B'})")
    print("-" * 50)
    
    if wb_type == "legacy":
        print("[INFO] This is a legacy Contract A workbook. To inspect, please use standard generate/post commands or convert it.")
        return
        
    # It's a business workbook!
    sheet_scope = [sheet] if sheet else None
    try:
        rows, warnings = ingest_business_workbook(filepath, sheet_scope=sheet_scope)
    except Exception as e:
        print(f"[ERROR] Failed to ingest business workbook: {e}")
        sys.exit(1)
        
    if warnings:
        print("Warnings/Parse Info:")
        for w in warnings:
            print(f"  - {w}")
        print("-" * 50)
        
    if as_plan:
        # Generate the plan
        p_list = [p.strip() for p in platforms_str.split(",") if p.strip()]
        try:
            config = AppConfig()
            plan = build_posting_plan(
                rows,
                p_list,
                limit_per_platform=limit,
                config=config,
                workbook_path=filepath,
            )
        except Exception as e:
            print(f"[ERROR] Failed to build posting plan: {e}")
            sys.exit(1)
            
        if as_json:
            import json
            print(json.dumps(plan.to_dict(), indent=2, ensure_ascii=False))
        else:
            print(f"POSTING PLAN (Limit: {limit} posts per platform)")
            print("=" * 60)
            
            for platform, candidates in plan.candidates.items():
                platform_display = platform.upper()
                print(f"\nPlatform: {platform_display}")
                print("-" * 30)
                if not candidates:
                    print("  No eligible candidates to post.")
                else:
                    for idx, c in enumerate(candidates):
                        print(f"  {idx + 1}. [Sheet: {c.sheet_name}, Row {c.row_idx}, ID: {c.id}]")
                        print(f"     Title: {c.title}")
                        print(f"     Image: {c.image_link}")
                        # Show a short snippet of draft content
                        draft_snippet = c.draft_content.replace('\n', ' ')
                        if len(draft_snippet) > 80:
                            draft_snippet = draft_snippet[:80] + "..."
                        print(f"     Draft: {draft_snippet}")
                print("-" * 30)
                
            if plan.skipped_summary:
                print("\nSkipped/Excluded rows details:")
                print("=" * 40)
                for skip_info in plan.skipped_summary:
                    print(f"  - {skip_info}")
    else:
        if as_json:
            # Print JSON output
            import json
            dict_rows = [r.to_dict() for r in rows]
            print(json.dumps(dict_rows, indent=2, ensure_ascii=False))
        else:
            # Print human-readable output
            print(f"Total category rows parsed: {len(rows)}")
            current_sheet = None
            for r in rows:
                if r.sheet_name != current_sheet:
                    current_sheet = r.sheet_name
                    print(f"\nSheet: [{current_sheet}]")
                    print("=" * 40)
                
                # Print row summary
                has_content_list = []
                for plat in ['linkedin', 'facebook', 'x', 'instagram', 'pinterest', 'threads', 'tiktok', 'youtube']:
                    draft_val = getattr(r, f"{plat}_draft")
                    if draft_val:
                        has_content_list.append(plat)
                        
                platforms_str_val = ", ".join(has_content_list) if has_content_list else "None"
                print(f"Row {r.row_idx} (ID: {r.id}): {r.title}")
                print(f"  Image Link: {r.image_link}")
                print(f"  Drafts for: {platforms_str_val}")
                if r.link_post_raw:
                    print("  Link Post State:")
                    for line in r.link_post_raw.splitlines():
                        print(f"    {line}")
                print("-" * 40)

def execute_write_result(
    sheet: str,
    row: int,
    platform: str,
    status: str,
    workbook: str = None,
    no_backup: bool = False
):
    from src.writeback import write_post_result
    config = AppConfig()
    filepath = workbook or config.excel_file
    backup_enabled = not no_backup and config.backup_enabled
    
    try:
        new_val = write_post_result(
            workbook_path=filepath,
            sheet_name=sheet,
            row_idx=row,
            platform=platform,
            status_value=status,
            backup_enabled=backup_enabled
        )
        print(f"[SUCCESS] Wrote post result to sheet '{sheet}', row {row}, platform '{platform}'.")
        print(f"New Link Post value:\n{new_val}")
    except Exception as e:
        print(f"[ERROR] Failed to write result: {e}")
        sys.exit(1)

def main():
    parser = argparse.ArgumentParser(description="Trillion News Auto Post System CLI")
    subparsers = parser.add_subparsers(dest="command", help="Available commands")
    
    # init command
    subparsers.add_parser("init", help="Initialize project folders and configuration files")
    
    # search command
    search_parser = subparsers.add_parser("search", help="Search and filter news")
    search_parser.add_argument("--keywords", default="keywords.txt", help="Path to keywords file")
    search_parser.add_argument("--limit", type=int, default=None, help="Max results to return")
    
    # generate command
    generate_parser = subparsers.add_parser("generate", help="Generate AI social posts")
    generate_parser.add_argument("--platform", default=None, help="Target social platform (e.g. linkedin)")
    generate_parser.add_argument("--limit", type=int, default=None, help="Max posts to generate")
    
    # run command
    run_parser = subparsers.add_parser("run", help="Run full search and generation draft pipeline")
    run_parser.add_argument("--keywords", default="keywords.txt", help="Path to keywords file")
    run_parser.add_argument("--platform", default=None, help="Target social platform")
    run_parser.add_argument("--limit", type=int, default=None, help="Max posts to generate/return")
    
    # post command
    post_parser = subparsers.add_parser("post", help="Assisted posting on social platforms")
    post_parser.add_argument("--id", type=int, required=True, help="Row ID from Excel to post")
    post_parser.add_argument("--platform", default=None, help="Target social platform (e.g. linkedin)")
    post_parser.add_argument("--sheet", default=None, help="Optional specific sheet/category name for Business Workbook")

    # inspect-workbook command
    inspect_parser = subparsers.add_parser("inspect-workbook", help="Inspect and validate a Business Workbook")
    inspect_parser.add_argument("--workbook", default=None, help="Path to the business workbook file")
    inspect_parser.add_argument("--sheet", default=None, help="Optional specific sheet/category to inspect")
    inspect_parser.add_argument("--json", action="store_true", help="Output in structured JSON format")
    inspect_parser.add_argument("--plan", action="store_true", help="Generate and inspect the posting plan (dry-run)")
    inspect_parser.add_argument("--limit", type=int, default=2, help="Posting limit per platform (default: 2)")
    inspect_parser.add_argument("--platforms", default="linkedin,facebook,x,instagram,pinterest,threads,tiktok,youtube", help="Comma-separated platforms to plan for")

    # write-result command
    write_parser = subparsers.add_parser("write-result", help="Manually write a posting result to a row")
    write_parser.add_argument("--sheet", required=True, help="Name of the sheet/category")
    write_parser.add_argument("--row", type=int, required=True, help="1-based row index in the sheet")
    write_parser.add_argument("--platform", required=True, help="Target platform name (e.g. linkedin)")
    write_parser.add_argument("--status", required=True, help="Status value to write (URL or tag)")
    write_parser.add_argument("--workbook", default=None, help="Path to the business workbook file")
    write_parser.add_argument("--no-backup", action="store_true", help="Disable timestamped workbook backup before write")

    # web command
    web_parser = subparsers.add_parser("web", help="Start the local Web UI operator surface")
    web_parser.add_argument("--port", type=int, default=8080, help="Port to run local web server on (default: 8080)")

    # schedule command
    schedule_parser = subparsers.add_parser("schedule", help="Manage local schedules and run history")
    schedule_subparsers = schedule_parser.add_subparsers(dest="schedule_command", help="Schedule subcommands")
    
    # schedule add
    add_parser = schedule_subparsers.add_parser("add", help="Add a new schedule")
    add_parser.add_argument("--name", required=True, help="Name of the schedule")
    add_parser.add_argument("--expression", required=True, help="Cron expression or interval (e.g. '*/15 * * * *' or 'every 1h')")
    add_parser.add_argument("--job-type", required=True, choices=["draft", "post"], help="Type of job (draft = harvest & generate, post = publishing)")
    add_parser.add_argument("--workbook", default=None, help="Custom workbook file path")
    add_parser.add_argument("--sheet", default=None, help="Specific sheet name to post from")
    add_parser.add_argument("--platforms", default=None, help="Comma-separated platforms to post")
    add_parser.add_argument("--limit", type=int, default=2, help="Posting limit per platform")
    
    # schedule list
    schedule_subparsers.add_parser("list", help="List all schedule definitions")
    
    # schedule toggle
    toggle_parser = schedule_subparsers.add_parser("toggle", help="Enable or disable a schedule")
    toggle_parser.add_argument("--name", required=True, help="Name of the schedule")
    toggle_parser.add_argument("--enable", action="store_true", help="Enable the schedule")
    toggle_parser.add_argument("--disable", action="store_true", help="Disable the schedule")
    
    # schedule run-now
    run_now_parser = schedule_subparsers.add_parser("run-now", help="Trigger a schedule run immediately")
    run_now_parser.add_argument("--name", required=True, help="Name of the schedule")
    
    # schedule trigger-due
    schedule_subparsers.add_parser("trigger-due", help="Trigger any due schedules immediately")
    
    # schedule history
    history_parser = schedule_subparsers.add_parser("history", help="List run history")
    history_parser.add_argument("--limit", type=int, default=20, help="Max run records to return")
    
    # schedule history-detail
    detail_parser = schedule_subparsers.add_parser("history-detail", help="Show detail of a run")
    detail_parser.add_argument("--run-id", required=True, help="UUID of the run")

    args = parser.parse_args()
    
    if args.command == "init":
        init_project()
    elif args.command == "search":
        execute_search(args.keywords, args.limit)
    elif args.command == "generate":
        execute_generate(args.platform, args.limit)
    elif args.command == "run":
        execute_run(args.keywords, args.platform, args.limit)
    elif args.command == "post":
        execute_post(args.id, args.platform, args.sheet)
    elif args.command == "inspect-workbook":
        execute_inspect_workbook(
            workbook_path=args.workbook,
            sheet=args.sheet,
            as_json=args.json,
            as_plan=args.plan,
            limit=args.limit,
            platforms_str=args.platforms
        )
    elif args.command == "write-result":
        execute_write_result(
            sheet=args.sheet,
            row=args.row,
            platform=args.platform,
            status=args.status,
            workbook=args.workbook,
            no_backup=args.no_backup
        )
    elif args.command == "web":
        from src.web_ui import start_web_server
        start_web_server(port=args.port)
    elif args.command == "schedule":
        config = AppConfig()
        db_path = os.path.join(config.output_dir, "scheduler.db")
        
        from src.scheduler import (
            create_schedule, update_schedule, list_schedules, 
            run_schedule_now, trigger_due_schedules, list_run_history,
            get_run_history_detail
        )
        
        if args.schedule_command == "add":
            try:
                create_schedule(
                    db_path=db_path,
                    name=args.name,
                    expression=args.expression,
                    job_type=args.job_type,
                    workbook_path=args.workbook,
                    sheet_name=args.sheet,
                    platforms=args.platforms,
                    post_limit=args.limit
                )
                print(f"[SUCCESS] Created schedule '{args.name}' with expression '{args.expression}'")
            except Exception as e:
                print(f"[ERROR] Failed to add schedule: {e}")
                sys.exit(1)
                
        elif args.schedule_command == "list":
            schedules = list_schedules(db_path)
            if not schedules:
                print("No schedules defined.")
                return
            print(f"{'ID':<4} {'Name':<20} {'Recurrence':<18} {'Job Type':<10} {'Enabled':<8} {'Next Run':<25}")
            print("-" * 90)
            for s in schedules:
                enabled_str = "Yes" if s["enabled"] == 1 else "No"
                next_run_str = s["next_run_at"] or "N/A"
                print(f"{s['id']:<4} {s['name']:<20} {s['expression']:<18} {s['job_type']:<10} {enabled_str:<8} {next_run_str:<25}")
                
        elif args.schedule_command == "toggle":
            if not args.enable and not args.disable:
                print("[ERROR] Please specify either --enable or --disable")
                sys.exit(1)
            enabled_val = 1 if args.enable else 0
            try:
                update_schedule(db_path, args.name, enabled=enabled_val)
                state_str = "enabled" if args.enable else "disabled"
                print(f"[SUCCESS] Schedule '{args.name}' has been {state_str}.")
            except Exception as e:
                print(f"[ERROR] Failed to toggle schedule: {e}")
                sys.exit(1)
                
        elif args.schedule_command == "run-now":
            schedules = list_schedules(db_path)
            sch_id = None
            for s in schedules:
                if s["name"] == args.name:
                    sch_id = s["id"]
                    break
            if sch_id is None:
                print(f"[ERROR] Schedule '{args.name}' not found.")
                sys.exit(1)
                
            run_id = run_schedule_now(db_path, sch_id, config)
            print(f"[SUCCESS] Launched run immediately for schedule '{args.name}' in background (Run ID: {run_id}).")
            
        elif args.schedule_command == "trigger-due":
            trigger_due_schedules(db_path, config)
            print("[INFO] Finished triggering due schedules.")
            
        elif args.schedule_command == "history":
            history = list_run_history(db_path, args.limit)
            if not history:
                print("No execution history.")
                return
            print(f"{'Run ID':<38} {'Schedule':<20} {'Status':<12} {'Started At':<25} {'Outcome Summary'}")
            print("-" * 120)
            for r in history:
                sch_name = r["schedule_name"] or "(Deleted)"
                summary_str = r["summary"] or r["error_message"] or ""
                print(f"{r['run_id']:<38} {sch_name:<20} {r['status'].upper():<12} {r['started_at']:<25} {summary_str}")
                
        elif args.schedule_command == "history-detail":
            detail = get_run_history_detail(db_path, args.run_id)
            if not detail:
                print(f"[ERROR] Run history with ID '{args.run_id}' not found.")
                sys.exit(1)
            print(f"Run ID:        {detail['run_id']}")
            print(f"Schedule Name: {detail['schedule_name'] or '(Deleted)'}")
            print(f"Trigger Type:  {detail['trigger_type']}")
            print(f"Status:        {detail['status'].upper()}")
            print(f"Started At:    {detail['started_at']}")
            print(f"Finished At:   {detail['finished_at'] or 'N/A'}")
            print(f"Summary:       {detail['summary'] or 'N/A'}")
            if detail['error_message']:
                print(f"Error Message: {detail['error_message']}")
            print("\nRow-Level Posting Details:")
            print("=" * 100)
            if not detail['details']:
                print("No details recorded.")
            else:
                print(f"{'Sheet':<15} {'Row':<5} {'Platform':<12} {'Status':<10} {'Message'}")
                print("-" * 100)
                for d in detail['details']:
                    print(f"{d['sheet_name']:<15} {d['row_id']:<5} {d['platform']:<12} {d['status'].upper():<10} {d['message']}")
        else:
            schedule_parser.print_help()
    else:
        parser.print_help()

if __name__ == "__main__":
    main()

