import datetime
import os
import shutil
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

GENERATED_POST_MARKER = "## Generated Post"


@dataclass
class PostDraftRecord:
    relative_path: str
    filename: str
    absolute_path: str
    modified_at: Optional[str]
    title_hint: Optional[str]
    workbook_link: Optional[Dict[str, object]]


@dataclass
class PostDraftDetail:
    relative_path: str
    filename: str
    absolute_path: str
    modified_at: Optional[str]
    raw_content: str
    generated_post: str
    preamble: str
    news_section: Optional[str]
    image_section: Optional[str]
    workbook_link: Optional[Dict[str, object]]


def resolve_post_dir(post_dir: Optional[str] = None, config=None) -> str:
    if post_dir:
        return os.path.abspath(post_dir)
    if config is not None:
        configured = getattr(config, "post_dir", None)
        if configured:
            return os.path.abspath(configured)
    return os.path.abspath(os.getenv("POST_DIR", os.path.join("output", "posts")))


def validate_relative_post_path(relative_path: str, post_dir: str) -> Tuple[str, str]:
    if not relative_path or not str(relative_path).strip():
        raise ValueError("relative_path is required")

    normalized = str(relative_path).replace("\\", "/").strip().lstrip("/")
    if not normalized:
        raise ValueError("relative_path is required")
    if normalized.startswith("/"):
        raise ValueError("relative_path must not be absolute")
    parts = [part for part in normalized.split("/") if part not in ("", ".")]
    if any(part == ".." for part in parts):
        raise ValueError("relative_path must not contain '..'")
    if not normalized.lower().endswith(".md"):
        raise ValueError("Only .md post drafts are supported")

    abs_post_dir = os.path.abspath(post_dir)
    abs_path = os.path.normpath(os.path.join(abs_post_dir, *parts))
    post_dir_prefix = abs_post_dir if abs_post_dir.endswith(os.sep) else abs_post_dir + os.sep
    if abs_path != abs_post_dir and not abs_path.startswith(post_dir_prefix):
        raise ValueError("relative_path resolves outside POST_DIR")

    return abs_path, "/".join(parts)


def split_post_sections(content: str) -> Tuple[str, str, Optional[str], Optional[str]]:
    """Return preamble (through marker), generated body, news block, image block."""
    marker = GENERATED_POST_MARKER
    idx = content.find(marker)
    if idx == -1:
        return "", content.strip(), None, None

    preamble = content[: idx + len(marker)]
    remainder = content[idx + len(marker) :]
    if remainder.startswith("\r\n"):
        remainder = remainder[2:]
    elif remainder.startswith("\n"):
        remainder = remainder[1:]

    generated_post = remainder.strip()
    news_section = _extract_section(content, "## News", ["## Image", marker])
    image_section = _extract_section(content, "## Image", [marker])
    return preamble, generated_post, news_section, image_section


def _extract_section(content: str, heading: str, stop_headings: List[str]) -> Optional[str]:
    idx = content.find(heading)
    if idx == -1:
        return None
    start = idx + len(heading)
    end = len(content)
    for stop in stop_headings:
        stop_idx = content.find(stop, start)
        if stop_idx != -1:
            end = min(end, stop_idx)
    section = content[start:end].strip()
    return section or None


def merge_generated_post_content(content: str, generated_post: str) -> str:
    preamble, _, _, _ = split_post_sections(content)
    body = (generated_post or "").strip("\n")
    if not preamble:
        if body:
            return f"{GENERATED_POST_MARKER}\n\n{body}\n"
        return f"{GENERATED_POST_MARKER}\n\n"
    if body:
        return f"{preamble.rstrip()}\n\n{body}\n"
    return f"{preamble.rstrip()}\n\n"


def _title_hint_from_filename(filename: str) -> Optional[str]:
    base = os.path.splitext(filename)[0]
    return base.replace("_", " ") if base else None


def _iso_mtime(path: str) -> Optional[str]:
    try:
        return datetime.datetime.fromtimestamp(os.path.getmtime(path)).isoformat()
    except OSError:
        return None


def post_relative_path_for_absolute(absolute_path: str, post_dir: str) -> Optional[str]:
    abs_post_dir = os.path.abspath(post_dir)
    abs_path = os.path.abspath(absolute_path)
    try:
        rel = os.path.relpath(abs_path, abs_post_dir)
    except ValueError:
        return None
    if rel.startswith(".."):
        return None
    return rel.replace("\\", "/")


def build_workbook_post_links(config) -> Dict[str, Dict[str, object]]:
    """Map post-draft relative paths to workbook row metadata."""
    workbook_path = getattr(config, "excel_file", None)
    post_dir = resolve_post_dir(config=config)
    if not workbook_path or not os.path.exists(workbook_path):
        return {}

    try:
        import openpyxl
    except ImportError:
        return {}

    from src.business_workbook import normalize_header, normalize_cell_content
    from src.draft_paths import looks_like_file_reference, resolve_post_draft_path

    links: Dict[str, Dict[str, object]] = {}
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            header_row = []
            for row in ws.iter_rows(max_row=1, values_only=True):
                header_row = row
                break
            if not header_row:
                continue

            col_map = {}
            for idx, val in enumerate(header_row):
                norm = normalize_header(val)
                if norm:
                    col_map[norm] = idx + 1

            linkedin_col = col_map.get("linkedin")
            title_col = col_map.get("trillion $ news title")
            id_col = col_map.get("#", 1)
            if not linkedin_col:
                continue

            linkedin_idx = linkedin_col - 1
            title_idx = title_col - 1 if title_col else None
            id_idx = id_col - 1

            row_idx = 1
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_idx += 1
                if linkedin_idx >= len(row):
                    continue
                raw_linkedin = normalize_cell_content(row[linkedin_idx])
                if not raw_linkedin or not looks_like_file_reference(raw_linkedin):
                    continue

                resolved = resolve_post_draft_path(
                    raw_linkedin,
                    workbook_path,
                    post_dir,
                )
                if not resolved:
                    continue

                rel_path = post_relative_path_for_absolute(resolved, post_dir)
                if not rel_path:
                    continue

                title = None
                if title_idx is not None and title_idx < len(row):
                    title = normalize_cell_content(row[title_idx])
                row_id = row[id_idx] if id_idx < len(row) else None

                links[rel_path] = {
                    "sheet_name": sheet_name,
                    "row_idx": row_idx,
                    "row_id": row_id,
                    "title": title,
                    "linkedin_draft_ref": raw_linkedin,
                }
    finally:
        wb.close()

    return links


def list_post_drafts(config, query: Optional[str] = None) -> List[PostDraftRecord]:
    post_dir = resolve_post_dir(config=config)
    if not os.path.isdir(post_dir):
        return []

    workbook_links = build_workbook_post_links(config)
    query_norm = (query or "").strip().casefold()
    records: List[PostDraftRecord] = []

    for entry in sorted(os.listdir(post_dir)):
        if not entry.lower().endswith(".md"):
            continue
        abs_path = os.path.join(post_dir, entry)
        if not os.path.isfile(abs_path):
            continue

        rel_path = entry
        title_hint = workbook_links.get(rel_path, {}).get("title") or _title_hint_from_filename(entry)
        if query_norm:
            haystack = f"{entry} {title_hint or ''}".casefold()
            if query_norm not in haystack:
                continue

        records.append(
            PostDraftRecord(
                relative_path=rel_path,
                filename=entry,
                absolute_path=abs_path,
                modified_at=_iso_mtime(abs_path),
                title_hint=title_hint,
                workbook_link=workbook_links.get(rel_path),
            )
        )

    return records


def read_post_draft(config, relative_path: str) -> PostDraftDetail:
    post_dir = resolve_post_dir(config=config)
    abs_path, rel_path = validate_relative_post_path(relative_path, post_dir)
    if not os.path.exists(abs_path):
        raise FileNotFoundError(f"Post draft '{rel_path}' does not exist")

    with open(abs_path, "r", encoding="utf-8") as f:
        raw_content = f.read()

    preamble, generated_post, news_section, image_section = split_post_sections(raw_content)
    workbook_links = build_workbook_post_links(config)

    return PostDraftDetail(
        relative_path=rel_path,
        filename=os.path.basename(abs_path),
        absolute_path=abs_path,
        modified_at=_iso_mtime(abs_path),
        raw_content=raw_content,
        generated_post=generated_post,
        preamble=preamble,
        news_section=news_section,
        image_section=image_section,
        workbook_link=workbook_links.get(rel_path),
    )


def save_post_draft(
    config,
    relative_path: str,
    generated_post: str,
    *,
    create_backup: bool = True,
) -> PostDraftDetail:
    post_dir = resolve_post_dir(config=config)
    abs_path, rel_path = validate_relative_post_path(relative_path, post_dir)
    if not os.path.exists(abs_path):
        raise FileNotFoundError(f"Post draft '{rel_path}' does not exist")

    with open(abs_path, "r", encoding="utf-8") as f:
        raw_content = f.read()

    updated_content = merge_generated_post_content(raw_content, generated_post)

    if create_backup:
        backup_path = f"{abs_path}.bak"
        try:
            shutil.copy2(abs_path, backup_path)
        except OSError as exc:
            raise PermissionError(f"Failed to create backup for '{rel_path}': {exc}") from exc

    try:
        with open(abs_path, "w", encoding="utf-8") as f:
            f.write(updated_content)
    except OSError as exc:
        raise PermissionError(f"Failed to write post draft '{rel_path}': {exc}") from exc

    return read_post_draft(config, rel_path)


def post_draft_record_to_api(record: PostDraftRecord) -> dict:
    return {
        "relative_path": record.relative_path,
        "filename": record.filename,
        "modified_at": record.modified_at,
        "title_hint": record.title_hint,
        "workbook_link": record.workbook_link,
    }


def post_draft_detail_to_api(detail: PostDraftDetail) -> dict:
    return {
        "relative_path": detail.relative_path,
        "filename": detail.filename,
        "modified_at": detail.modified_at,
        "generated_post": detail.generated_post,
        "news_section": detail.news_section,
        "image_section": detail.image_section,
        "workbook_link": detail.workbook_link,
        "raw_content": detail.raw_content,
    }


def linkedin_draft_file_for_row(
    linkedin_draft_ref: Optional[str],
    config,
) -> Optional[str]:
    """Return API-relative post path for a workbook LinkedIn draft cell reference."""
    if not linkedin_draft_ref:
        return None

    from src.draft_paths import looks_like_file_reference, resolve_post_draft_path

    if not looks_like_file_reference(linkedin_draft_ref):
        return None

    post_dir = resolve_post_dir(config=config)
    workbook_path = getattr(config, "excel_file", None)
    resolved = resolve_post_draft_path(linkedin_draft_ref, workbook_path, post_dir)
    if not resolved:
        return None
    return post_relative_path_for_absolute(resolved, post_dir)