import os
import shutil
from datetime import datetime
from typing import Optional, List, Tuple
import openpyxl

from src.models import BusinessWorkbookRow, PlatformPostState, LinkPostDocument
from src.business_workbook import normalize_header
from src.posting_core import (
    canonicalize_platform,
    parse_link_post_document,
    serialize_link_post_document,
    PLATFORM_CANONICAL
)

class WorkbookLockedError(PermissionError):
    """Raised when the workbook is locked by another program."""
    pass

class BackupFailureError(IOError):
    """Raised when the workbook backup fails."""
    pass


def is_workbook_locked(filepath: str) -> bool:
    """
    Check if the workbook is locked by attempting to open it in read-write mode.
    """
    if not os.path.exists(filepath):
        return False
    try:
        # Open in read-write mode without truncating
        with open(filepath, "r+"):
            return False
    except PermissionError:
        return True
    except Exception:
        return True


def create_workbook_backup(filepath: str) -> str:
    """
    Create a timestamped copy of the workbook in the same directory.
    Returns the path to the backup file.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Workbook file '{filepath}' does not exist.")
        
    dir_name = os.path.dirname(filepath)
    base_name = os.path.basename(filepath)
    name, ext = os.path.splitext(base_name)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_name = f"{name}.{timestamp}{ext}"
    backup_path = os.path.join(dir_name, backup_name)
    
    try:
        shutil.copy2(filepath, backup_path)
        return backup_path
    except Exception as e:
        raise BackupFailureError(f"Failed to create workbook backup: {e}")


def evaluate_retry_disposition(link_post_raw: Optional[str], platform: str) -> bool:
    """
    Evaluate whether a prior posting attempt for a platform can be retried.
    Returns True if retryable (status is error, pending, login-required, or missing).
    Returns False if already posted successfully or explicitly skipped.
    """
    try:
        platform_key = canonicalize_platform(platform)
    except ValueError:
        return False
        
    if not link_post_raw:
        return True
        
    try:
        doc = parse_link_post_document(link_post_raw)
    except Exception:
        # If Link Post is malformed, we return False to be safe (cannot safely evaluate)
        return False
        
    if platform_key not in doc.states:
        return True
        
    state = doc.states[platform_key]
    return state.status_type == "retryable"


def write_post_result(
    workbook_path: str,
    sheet_name: str,
    row_idx: int,
    platform: str,
    status_value: str,
    backup_enabled: bool = True
) -> str:
    """
    Write the posting result of a platform to a specific row in the business workbook.
    Mutates only the target platform's line in the Link Post column.
    
    Args:
        workbook_path: Path to the Excel workbook.
        sheet_name: Name of the sheet to modify.
        row_idx: 1-based row index in the sheet.
        platform: Platform name (variant).
        status_value: The value to set (URL, [posted-no-link], [skip], [error] reason, etc.).
        backup_enabled: Whether to create a backup file before writing.
        
    Returns:
        The serialized new value of the Link Post cell.
    """
    if not os.path.exists(workbook_path):
        raise FileNotFoundError(f"Workbook file '{workbook_path}' not found.")
        
    if is_workbook_locked(workbook_path):
        raise WorkbookLockedError(
            f"Cannot write to workbook. File '{workbook_path}' is currently open/locked by another application."
        )
        
    platform_key = canonicalize_platform(platform)
    canonical_display = PLATFORM_CANONICAL[platform_key]
    
    # 1. Load workbook
    wb = openpyxl.load_workbook(workbook_path)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
            
        ws = wb[sheet_name]
        
        if row_idx < 2 or row_idx > ws.max_row:
            raise ValueError(f"Row index {row_idx} is out of bounds for sheet '{sheet_name}'.")
            
        # 2. Locate Link Post column
        # Read header row
        header_row = []
        for r in ws.iter_rows(max_row=1, values_only=True):
            header_row = r
            break
            
        col_idx = -1
        for idx, cell_val in enumerate(header_row):
            if cell_val is not None:
                norm = normalize_header(cell_val)
                if norm == "link post":
                    col_idx = idx + 1
                    break
                    
        # If missing, we will create the Link Post column at the end
        if col_idx == -1:
            col_idx = len(header_row) + 1
            ws.cell(row=1, column=col_idx).value = "Link Post"
            
        # 3. Read current value and parse Link Post document
        current_val = ws.cell(row=row_idx, column=col_idx).value
        # If parsing fails, raise ValueError to identify sheet and row context
        try:
            link_post_doc = parse_link_post_document(current_val)
        except Exception as e:
            raise ValueError(
                f"Cannot mutate Link Post at Row {row_idx} in sheet '{sheet_name}': cell is malformed. Details: {e}"
            )
            
        # 4. Prepare mutation
        val_lower = status_value.lower()
        if val_lower.startswith("http://") or val_lower.startswith("https://"):
            status_type = "success"
        elif val_lower.startswith("[posted-no-link]"):
            status_type = "success"
        elif val_lower.startswith("[skip]"):
            status_type = "skip"
        elif val_lower.startswith("[pending]") or val_lower.startswith("[error]") or val_lower.startswith("[login-required]"):
            status_type = "retryable"
        else:
            raise ValueError(
                f"Invalid status value: '{status_value}'. "
                f"Must be a URL or start with [posted-no-link], [skip], [pending], [error], [login-required]."
            )
            
        link_post_doc.states[platform_key] = PlatformPostState(
            platform=canonical_display,
            status_type=status_type,
            raw_value=status_value
        )
        
        new_raw_val = serialize_link_post_document(link_post_doc)
        
        # 5. Create backup if enabled
        if backup_enabled:
            # We must close the workbook handle before backing up to avoid lock warnings on some OSs
            wb.close()
            create_workbook_backup(workbook_path)
            # Re-open workbook
            wb = openpyxl.load_workbook(workbook_path)
            ws = wb[sheet_name]
            
        # 6. Write mutated cell value and save
        ws.cell(row=row_idx, column=col_idx).value = new_raw_val
        
        try:
            wb.save(workbook_path)
        except PermissionError:
            raise WorkbookLockedError(
                f"Failed to save workbook. File '{workbook_path}' is open/locked by another application."
            )
            
        return new_raw_val
        
    finally:
        wb.close()
