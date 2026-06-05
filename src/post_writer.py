import os
import re
from src.models import NewsItem
from src.config import AppConfig

# Load openpyxl safely
try:
    import openpyxl
except ImportError:
    openpyxl = None

def extract_top_hashtags(content: str) -> str:
    if not content:
        return ""
    lines = [line.strip() for line in content.split("\n") if line.strip()]
    if not lines:
        return ""
    first_line = lines[0]
    tags = re.findall(r'#\w+', first_line)
    return " ".join(tags[:5])

def write_post_file(item: NewsItem, content: str, config: AppConfig) -> str:
    os.makedirs(config.post_dir, exist_ok=True)
    
    padded_id = f"{item.id:03d}" if isinstance(item.id, int) else str(item.id)
    filename = f"{item.found_date}_{padded_id}_{item.platform}.md"
    filepath = os.path.join(config.post_dir, filename)
    
    # Calculate relative image path
    img_relative = ""
    if item.image_file:
        img_relative = f"../Ảnh Trillion $ news/{item.image_file}"
        
    markdown_content = f"""# Post {padded_id} — {item.platform.capitalize() if item.platform else ""}

## News

Title: {item.title}
Source: {item.source or "N/A"}
URL: {item.url}

## Image

{img_relative}

## Generated Post

{content}
"""
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(markdown_content)
        
    print(f"[SUCCESS] Saved post file: {filepath}")
    return filepath

def update_excel_row_with_post(item_id: int, post_file_path: str = None, top_hashtags: str = None, status: str = None, config: AppConfig = None, notes: str = None) -> bool:
    if openpyxl is None:
        print("[ERROR] openpyxl is not installed. Excel update skipped.")
        return False
        
    filepath = config.excel_file
    if not os.path.exists(filepath):
        print(f"[ERROR] Excel file '{filepath}' does not exist.")
        return False
        
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        
        row_found = False
        for r in range(2, ws.max_row + 1):
            id_val = ws.cell(row=r, column=1).value
            try:
                current_id = int(id_val)
            except (ValueError, TypeError):
                current_id = id_val
                
            if current_id == item_id:
                if status is not None:
                    ws.cell(row=r, column=13).value = status
                if post_file_path is not None:
                    ws.cell(row=r, column=12).value = post_file_path
                if top_hashtags is not None:
                    ws.cell(row=r, column=11).value = top_hashtags
                if notes is not None:
                    ws.cell(row=r, column=14).value = notes
                row_found = True
                break
                
        if row_found:
            wb.save(filepath)
            status_str = f"status '{status}'" if status is not None else "fields"
            print(f"[SUCCESS] Updated Excel row ID {item_id} with {status_str}.")
            return True
        else:
            print(f"[WARNING] Excel row ID {item_id} not found.")
            return False
            
    except PermissionError:
        print(f"[ERROR] Cannot write to Excel workbook. File '{filepath}' is currently open/locked.")
        return False
    except Exception as e:
        print(f"[ERROR] Failed to update Excel row: {e}")
        return False
