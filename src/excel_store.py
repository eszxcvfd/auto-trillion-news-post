import os
try:
    import openpyxl
except ImportError:
    openpyxl = None
from datetime import datetime
from src.models import NewsItem
from src.config import AppConfig

HEADERS = [
    "ID", "Found Date", "Keyword", "Title", "Source", "URL", 
    "Snippet", "Published Text", "Image File", "Platform", 
    "Top Hashtags", "Generated Post File", "Status", "Notes"
]

def init_excel_file(filepath: str):
    if openpyxl is None:
        print("[ERROR] openpyxl is not installed. Scaffolding excel file is skipped. Run pip install -r requirements.txt")
        raise PermissionError("openpyxl is not installed")

    # Ensure directory exists
    dir_name = os.path.dirname(filepath)
    if dir_name:
        os.makedirs(dir_name, exist_ok=True)
    
    if os.path.exists(filepath):
        return
        
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Trillion News"
        ws.append(HEADERS)
        wb.save(filepath)
        print(f"[SUCCESS] Initialized new Excel file: {filepath}")
    except PermissionError:
        print(f"[ERROR] Excel file {filepath} is currently locked or open in another application. Please close it.")
        raise

def get_next_id(sheet) -> int:
    if sheet.max_row <= 1:
        return 1
    # Check the last row's first column (ID)
    last_val = sheet.cell(row=sheet.max_row, column=1).value
    try:
        return int(last_val) + 1
    except (ValueError, TypeError):
        return sheet.max_row

import re

def slugify(text: str) -> str:
    text = text.lower()
    text = re.sub(r'[^a-z0-9_]', '_', text)
    text = re.sub(r'_+', '_', text)
    return text.strip('_')

def save_news_to_excel(items: list[NewsItem], config: AppConfig) -> list[NewsItem]:
    filepath = config.excel_file
    
    if openpyxl is None:
        print("[ERROR] openpyxl is not installed. Please run pip install -r requirements.txt to install it.")
        return []
        
    try:
        init_excel_file(filepath)
    except PermissionError:
        return []
        
    saved_items = []
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        
        # Collect existing URLs to prevent duplicating entries already in Excel
        existing_urls = set()
        for r in range(2, ws.max_row + 1):
            url_val = ws.cell(row=r, column=6).value
            if url_val:
                existing_urls.add(url_val.strip().lower())
                
        for item in items:
            if item.url and item.url.strip().lower() in existing_urls:
                # Delete temp image if duplicate
                if item.image_file and item.image_file.startswith("temp_"):
                    temp_path = os.path.join(config.image_dir, item.image_file)
                    if os.path.exists(temp_path):
                        try:
                            os.remove(temp_path)
                        except Exception:
                            pass
                print(f"[INFO] Skipping already saved URL in Excel: {item.url}")
                continue
                
            next_id = get_next_id(ws)
            item.id = next_id
            if not item.found_date:
                item.found_date = datetime.now().strftime("%Y-%m-%d")
            
            # Handle image renaming
            if item.image_file and item.image_file.startswith("temp_"):
                old_filename = item.image_file
                old_path = os.path.join(config.image_dir, old_filename)
                
                ext = os.path.splitext(old_filename)[1]
                slug_kw = slugify(item.keyword or "news")
                padded_id = f"{item.id:03d}" if isinstance(item.id, int) else str(item.id)
                new_filename = f"{item.found_date}_{padded_id}_{slug_kw}{ext}"
                new_path = os.path.join(config.image_dir, new_filename)
                
                if os.path.exists(old_path):
                    try:
                        os.rename(old_path, new_path)
                        item.image_file = new_filename
                    except Exception as e:
                        print(f"[WARNING] Failed to rename temp screenshot file: {e}")
                else:
                    item.image_file = None
            
            item.status = "new"
            item.platform = config.default_platform
            
            row_data = [
                item.id,
                item.found_date,
                item.keyword,
                item.title,
                item.source,
                item.url,
                item.snippet,
                item.published_text,
                item.image_file,
                item.platform,
                item.top_hashtags,
                item.generated_post_file,
                item.status,
                item.notes
            ]
            ws.append(row_data)
            saved_items.append(item)
            
        wb.save(filepath)
        if saved_items:
            print(f"[SUCCESS] Appended {len(saved_items)} new articles to Excel workbook.")
    except PermissionError:
        print(f"[ERROR] Cannot write to Excel workbook. The file '{filepath}' is currently open/locked by another application.")
        print("[INFO] Please close the file and rerun the command.")
        return []
        
    return saved_items
